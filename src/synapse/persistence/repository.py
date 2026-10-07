"""SQLite behind the CLI: ingest, search, count, and the RAG plumbing.

Owns the engine (one module-global, WAL + a patient timeout so the agents'
threads wait instead of racing) and every query. `models` owns the shape; this
module owns the access. `to_row` is the `--json` contract.

The database is the source of truth. `ingest_*` is how rows get in (a generic
xlsx/csv reader, or discovery records from `enrichment`); there is no build step
and no `--json`-visible difference between an ingested row and a discovered one.
"""

from __future__ import annotations

import ast
import csv
import re
from collections.abc import Iterable, Mapping, Sequence
from datetime import date
from pathlib import Path
from typing import Any, cast

from sqlalchemy import Engine, delete, func
from sqlalchemy.orm import selectinload
from sqlmodel import Session, SQLModel, create_engine, select

from ..config import DB_PATH, Settings
from ..shared.text import clean as _clean
from ..shared.text import norm_title
from .models import (
    ALL_MODELS,
    Author,
    Chunk,
    Citation,
    File,
    Paper,
    PaperAuthor,
    PaperTopic,
    Topic,
    Venue,
)

_engine: Engine | None = None


# --- engine -----------------------------------------------------------------


def _create_virtual_tables(engine: Engine) -> None:
    """The one table SQLModel cannot declare: FTS5, over `chunk.text`."""
    with engine.begin() as conn:
        conn.exec_driver_sql(
            "CREATE VIRTUAL TABLE IF NOT EXISTS chunk_fts USING fts5(text, tokenize='unicode61')"
        )


def connect(url: str | None = None) -> Engine:
    """The engine, created on first use, with the schema in place. Idempotent."""
    global _engine
    if _engine is None or url is not None:
        target = url or Settings().database_url
        # `data/` ships with the repo but `SYNAPSE_ROOT` can point anywhere, and
        # SQLite will not create a missing parent directory for you
        if target.startswith("sqlite:///"):
            DB_PATH.parent.mkdir(parents=True, exist_ok=True)
        engine = create_engine(
            target,
            # four agent threads each open their own session; without these the
            # second one to write gets "database is locked" rather than waiting
            connect_args={"timeout": 30, "check_same_thread": False}
            if target.startswith("sqlite")
            else {},
        )
        with engine.connect() as connection:
            if target.startswith("sqlite"):
                # WAL lets a reader and a writer coexist, which the agents need
                connection.exec_driver_sql("PRAGMA journal_mode=WAL")
        SQLModel.metadata.create_all(engine)
        if target.startswith("sqlite"):
            _create_virtual_tables(engine)
        _engine = engine
    return _engine


# --- ingest -----------------------------------------------------------------

# What a column header can be called, mapped to the one canonical name. First
# spelling wins, so a sheet with both `date` and `year` keeps whichever comes
# first rather than silently overwriting one with the other.
FIELD_ALIASES: dict[str, str] = {
    "title": "title",
    "name": "title",
    "year": "year",
    "date": "year",
    "publication_year": "year",
    "published": "year",
    "doi": "doi",
    "url": "url",
    "link": "url",
    "landing_page": "url",
    "abstract": "abstract",
    "summary": "abstract",
    "authors": "authors",
    "author": "authors",
    "venue": "venue",
    "journal": "venue",
    "container": "venue",
    "topics": "topics",
    "topic": "topics",
    "keywords": "topics",
    "tags": "topics",
    "kind": "kind",
    "type": "kind",
    "item_type": "kind",
    "source": "source",
}


def normalize_keys(record: Mapping[Any, Any]) -> dict[str, Any]:
    """A row from any spreadsheet -> the canonical field names, first spelling wins."""
    out: dict[str, Any] = {}
    for raw_key, value in record.items():
        key = re.sub(r"[^a-z0-9]+", "_", str(raw_key).strip().lower()).strip("_")
        canonical = FIELD_ALIASES.get(key)
        if canonical and canonical not in out:
            out[canonical] = value
    return out


def _split_list(value: object) -> list[str]:
    if value is None:
        return []
    if isinstance(value, (list, tuple, set)):
        return [item for item in (_clean(v) for v in value) if item]
    return [item for item in (_clean(v) for v in re.split(r"[;,]", str(value))) if item]


def _split_authors(value: object) -> list[str]:
    """Authors split on `;` only - a comma is how half the world writes one name."""
    if value is None:
        return []
    if isinstance(value, (list, tuple, set)):
        return [item for item in (_clean(v) for v in value) if item]
    text = _clean(value)
    if text.startswith("["):
        try:
            parsed = ast.literal_eval(text)
        except ValueError, SyntaxError:
            parsed = None
        if isinstance(parsed, list):
            return [item for item in (_clean(v) for v in parsed) if item]
    return [item for item in (_clean(v) for v in text.split(";")) if item]


def _year_of(value: object) -> int:
    try:
        return int(float(str(value)))  # type: ignore[arg-type]
    except TypeError, ValueError:
        return 0


def _bool_of(value: object) -> bool:
    return _clean(value).upper() in {"TRUE", "YES", "1", "Y"}


def normalize_record(record: Mapping[Any, Any]) -> dict[str, Any]:
    """Any source row -> the one dict `upsert_paper` understands.

    Spreadsheets go through `normalize_keys`; discovery records are already
    canonical, and both land here so there is a single ingest path.
    """
    if "title" not in record:
        record = normalize_keys(record)
    topics = record.get("topics")
    parsed_topics: list[dict[str, str]] = []
    if isinstance(topics, (list, tuple)):
        for item in topics:
            if isinstance(item, Mapping):
                parsed_topics.append(
                    {"name": _clean(item.get("name")), "category": _clean(item.get("category"))}
                )
            else:
                parsed_topics.append({"name": _clean(item), "category": ""})
    else:
        parsed_topics = [{"name": name, "category": ""} for name in _split_list(topics)]
    authors = record.get("authors")
    meta = record.get("meta")
    extraction = record.get("extraction")
    return {
        "title": _clean(record.get("title")),
        "year": _year_of(record.get("year")),
        "doi": _clean(record.get("doi")),
        "url": _clean(record.get("url")),
        "abstract": _clean(record.get("abstract")),
        "abstract_source": _clean(record.get("abstract_source")),
        "venue": _clean(record.get("venue")),
        "kind": _clean(record.get("kind")) or "article",
        "source": _clean(record.get("source")),
        "topics": [t for t in parsed_topics if t["name"]],
        "authors": (
            [_clean(a) for a in authors] if isinstance(authors, list) else _split_authors(authors)
        ),
        "inferred": record.get("inferred"),
        "needs_review": record.get("needs_review"),
        "meta": meta if isinstance(meta, dict) else {},
        "extraction": extraction if isinstance(extraction, dict) else {},
    }


def get_or_create(session: Session, model: type[SQLModel], **keys: Any) -> Any:
    """The row matching `keys`, created if absent. One query, no race on re-import."""
    found = session.exec(select(model).filter_by(**keys)).first()
    if found is None:
        found = model(**keys)
        session.add(found)
    return found


def _apply_scalars(paper: Paper, rec: Mapping[str, Any]) -> None:
    """Fill non-empty scalars without blanking an existing value.

    Discovery and seeding only ever add information; a blank from a thinner
    source must not erase a curated value. The flags are explicit, so they do
    overwrite when the caller passes them.
    """
    if rec.get("kind"):
        paper.kind = rec["kind"]
    if rec.get("doi") and not paper.doi:
        paper.doi = rec["doi"]
    if rec.get("url") and not paper.url:
        paper.url = rec["url"]
    if rec.get("abstract"):
        current = paper.abstract
        # a full abstract beats a short fragment; a fragment never overwrites a full one
        if not current or len(rec["abstract"]) > len(current):
            paper.abstract = rec["abstract"]
    if rec.get("abstract_source") and not paper.abstract_source:
        paper.abstract_source = rec["abstract_source"]
    if rec.get("source"):
        paper.source = rec["source"]
    if rec.get("year"):
        paper.year = rec["year"]
    if rec.get("inferred") is not None:
        paper.inferred = _bool_of(rec["inferred"])
    if rec.get("needs_review") is not None:
        paper.needs_review = _bool_of(rec["needs_review"])
    if rec.get("meta"):
        paper.meta = {**paper.meta, **rec["meta"]}
    if rec.get("extraction"):
        paper.extraction = rec["extraction"]


def upsert_paper(session: Session, record: Mapping[Any, Any], *, source: str = "") -> Paper | None:
    """Create or update one Paper and its edges. Returns None for a blank title.

    Dedups on normalized title first, then on DOI, because the same work arrives
    spelled two ways ("Grey Wolf Optimizer: a new metaheuristic" vs "Grey wolf
    optimizer"). The title is never rewritten on a match, so the unique index
    cannot blow up on a near-duplicate.
    """
    rec = normalize_record(record)
    title = rec["title"]
    if not title:
        return None
    if source and not rec["source"]:
        rec["source"] = source

    paper = session.exec(select(Paper).where(Paper.norm_title == norm_title(title))).first()
    if paper is None and rec["doi"]:
        paper = session.exec(select(Paper).where(Paper.doi == rec["doi"])).first()
    if paper is None:
        paper = Paper(
            title=title,
            norm_title=norm_title(title),
            retrieved_at=date.today().isoformat(),
        )
        session.add(paper)
    _apply_scalars(paper, rec)

    venue = rec["venue"]
    if venue and paper.venue_id is None:
        paper.venue = get_or_create(session, Venue, name=venue)

    # by name, not by id: a freshly created author is still pending and has no id,
    # so `if author.id in known` would match every author and keep only one
    known = {author.name for author in paper.authors}
    for author_name in rec["authors"]:
        if author_name and author_name not in known:
            known.add(author_name)
            paper.authors.append(get_or_create(session, Author, name=author_name))

    have = {topic.name for topic in paper.topics}
    for topic in rec["topics"]:
        name = topic["name"]
        if not name or name in have:
            continue
        have.add(name)
        obj = get_or_create(session, Topic, name=name)
        if topic["category"] and not obj.category:
            obj.category = topic["category"]
        paper.topics.append(obj)
    return paper


def ingest_records(records: Iterable[Mapping[Any, Any]], *, source: str = "") -> int:
    """Upsert many records in one transaction. Returns how many had a title."""
    engine = connect()
    with Session(engine) as session:
        written = sum(
            upsert_paper(session, record, source=source) is not None for record in records
        )
        session.commit()
    return written


def ingest_file(path: Path) -> int:
    """Load an xlsx or csv of papers. Any header spelling in `FIELD_ALIASES` works."""
    return ingest_records(_read_records(path), source="ingest")


def _read_records(path: Path) -> list[dict[str, Any]]:
    """The sheet or the delimited file as rows, header row first."""
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return _read_sheet(path)
    if suffix in {".csv", ".tsv"}:
        with path.open(encoding="utf-8-sig", newline="") as handle:
            delimiter = "\t" if suffix == ".tsv" else ","
            return list(csv.DictReader(handle, delimiter=delimiter))
    raise ValueError(f"unsupported file type: {path.suffix!r} (use .xlsx or .csv)")


def _read_sheet(path: Path) -> list[dict[str, Any]]:
    """One openpyxl sheet as rows: the header row's cells become the keys."""
    from openpyxl import load_workbook  # lazy, so `synapse --help` never pays for it

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active
        iterator = rows.iter_rows(values_only=True) if rows is not None else iter(())
        header = ["" if value is None else str(value) for value in next(iterator, ())]
        return [dict(zip(header, values, strict=False)) for values in iterator]
    finally:
        workbook.close()


# --- reading ----------------------------------------------------------------

_EAGER = (
    selectinload(cast("Any", Paper.topics)),
    selectinload(cast("Any", Paper.venue)),
    selectinload(cast("Any", Paper.authors)),
    selectinload(cast("Any", Paper.files)),
)


def to_row(paper: Paper, *, include_meta: bool = False) -> dict[str, Any]:
    """The `--json` contract. Explicit, so it does not leak a new column by accident."""
    venue = paper.venue
    row: dict[str, Any] = {
        "id": paper.id,
        "kind": paper.kind,
        "title": paper.title,
        "year": paper.year,
        "doi": paper.doi,
        "url": paper.url,
        "abstract": paper.abstract,
        "authors": sorted(author.name for author in paper.authors),
        "venue": "" if venue is None else venue.name,
        "topics": sorted(
            ({"name": topic.name, "category": topic.category} for topic in paper.topics),
            key=lambda topic: topic["name"],
        ),
        "files": sorted(
            (
                {
                    "kind": file.kind,
                    "path": file.path,
                    "url": file.url,
                    "sha256": file.sha256,
                    "status": file.status,
                }
                for file in paper.files
            ),
            key=lambda file: file["path"],
        ),
        "extraction": paper.extraction,
        "source": paper.source,
        "retrieved_at": paper.retrieved_at,
        "inferred": paper.inferred,
        "needs_review": paper.needs_review,
    }
    if include_meta:
        row["meta"] = paper.meta
    return row


def all_rows(*, include_meta: bool = False) -> list[dict[str, Any]]:
    """Every paper as `to_row` dicts, eagerly loaded. One query, session-safe."""
    engine = connect()
    with Session(engine) as session:
        statement = select(Paper).options(*_EAGER)
        return [to_row(paper, include_meta=include_meta) for paper in session.exec(statement)]


def get_row(paper_id: int, *, include_meta: bool = False) -> dict[str, Any] | None:
    """One paper by id, or None. The `search` contract for a single known row."""
    engine = connect()
    with Session(engine) as session:
        paper = session.get(Paper, paper_id)
        return None if paper is None else to_row(paper, include_meta=include_meta)


def search(
    query: str = "",
    *,
    year_from: int | None = None,
    year_to: int | None = None,
    author: str = "",
    venue: str = "",
    doi: str = "",
    kind: str = "",
    topic: str = "",
    has_pdf: bool = False,
    source: str = "",
    needs_review: bool | None = None,
    exact: bool = False,
    sort: str = "year",
    limit: int = 20,
    offset: int = 0,
    include_meta: bool = False,
) -> list[dict[str, Any]]:
    """Filter the collection. Empty `query` browses rather than searches.

    Every filter is a SQL predicate in the same statement as the `LIMIT`, and the
    many-valued ones (author/topic/venue) go through a subquery rather than a join
    so one paper with three topics is not returned three times.
    """
    engine = connect()
    statement = select(Paper)
    if query:
        statement = statement.where(
            Paper.norm_title == norm_title(query)
            if exact
            else cast("Any", Paper.norm_title).contains(norm_title(query))
        )
    if year_from:
        statement = statement.where(Paper.year >= year_from)
    if year_to:
        statement = statement.where(Paper.year <= year_to)
    if doi:
        statement = statement.where(Paper.doi == doi)
    if kind:
        statement = statement.where(Paper.kind == kind)
    if source:
        statement = statement.where(Paper.source == source)
    if needs_review is not None:
        statement = statement.where(Paper.needs_review == needs_review)
    if author:
        written = (
            select(PaperAuthor.paper_id)
            .join(Author, Author.id == PaperAuthor.author_id)  # type: ignore[arg-type]
            .where(cast("Any", Author.name).contains(author))
        )
        statement = statement.where(cast("Any", Paper.id).in_(written))
    if topic:
        tagged = (
            select(PaperTopic.paper_id)
            .join(Topic, Topic.id == PaperTopic.topic_id)  # type: ignore[arg-type]
            .where(cast("Any", Topic.name).contains(topic))
        )
        statement = statement.where(cast("Any", Paper.id).in_(tagged))
    if venue:
        named = select(Venue.id).where(cast("Any", Venue.name).contains(venue))
        statement = statement.where(cast("Any", Paper.venue_id).in_(named))
    if has_pdf:
        present = select(File.paper_id).where(cast("Any", File.status).in_(("ok", "cached")))
        statement = statement.where(cast("Any", Paper.id).in_(present))

    order = cast("Any", Paper.title) if sort == "title" else cast("Any", Paper.year).desc()
    statement = statement.order_by(order).offset(offset).limit(limit)  # type: ignore[arg-type]
    statement = statement.options(*_EAGER)  # type: ignore[arg-type]
    with Session(engine) as session:
        return [to_row(paper, include_meta=include_meta) for paper in session.exec(statement).all()]


def frame() -> list[dict[str, object]]:
    """Every paper as flat rows for the charts: one row per paper-topic pair."""
    engine = connect()
    sql = (
        "SELECT p.year AS year, p.kind AS kind, "
        "COALESCE(t.name, '') AS topic, COALESCE(v.name, '') AS venue "
        "FROM paper p "
        "LEFT JOIN venue v ON v.id = p.venue_id "
        "LEFT JOIN paper_topic pt ON pt.paper_id = p.id "
        "LEFT JOIN topic t ON t.id = pt.topic_id"
    )
    with engine.connect() as connection:
        return [dict(row) for row in connection.exec_driver_sql(sql).mappings().all()]


def stats() -> dict[str, int]:
    """Row and category counts."""
    engine = connect()
    with Session(engine) as session:
        counts = {
            model.__name__: int(
                session.scalar(select(func.count()).select_from(model)) or 0  # type: ignore[arg-type]
            )
            for model in ALL_MODELS
        }
        categories = session.exec(
            select(Paper.kind, func.count()).group_by(Paper.kind).order_by(func.count().desc())
        ).all()
    return {**counts, **{f"kind: {kind or '(blank)'}": count for kind, count in categories}}


# --- files ------------------------------------------------------------------


def record_file(session: Session, paper_id: int, fields: Mapping[str, Any]) -> File | None:
    """Upsert one retrieved artifact for a paper. Keyed by (paper, path)."""
    path = _clean(fields.get("path"))
    if not path:
        return None
    file = session.exec(select(File).where(File.paper_id == paper_id, File.path == path)).first()
    if file is None:
        file = File(paper_id=paper_id, path=path)
        session.add(file)
    for field in ("kind", "url", "source", "sha256", "status", "note"):
        value = _clean(fields.get(field))
        if value:
            setattr(file, field, value)
    file.size_bytes = int(fields.get("size_bytes") or 0)
    file.retrieved_at = _clean(fields.get("retrieved_at")) or date.today().isoformat()
    return file


# --- citations --------------------------------------------------------------


def link_citation(session: Session, citing_id: int, cited_id: int) -> bool:
    """Record that one paper cites another. False for a self-cite or a duplicate."""
    if not citing_id or not cited_id or citing_id == cited_id:
        return False
    existing = session.exec(
        select(Citation).where(Citation.citing_id == citing_id, Citation.cited_id == cited_id)
    ).first()
    if existing is not None:
        return False
    session.add(Citation(citing_id=citing_id, cited_id=cited_id))
    return True


def _brief(paper: Paper) -> dict[str, Any]:
    return {
        "id": paper.id,
        "kind": paper.kind,
        "title": paper.title,
        "year": paper.year,
        "doi": paper.doi,
    }


def _papers(ids: Sequence[int]) -> list[dict[str, Any]]:
    if not ids:
        return []
    engine = connect()
    with Session(engine) as session:
        statement = select(Paper).where(cast("Any", Paper.id).in_(list(ids)))
        return [_brief(paper) for paper in session.exec(statement)]


def references_of(paper_id: int) -> list[dict[str, Any]]:
    """What this paper cites, newest first."""
    engine = connect()
    with Session(engine) as session:
        ids = [
            int(cited)
            for cited in session.exec(
                select(Citation.cited_id).where(Citation.citing_id == paper_id)
            )
            if cited is not None
        ]
    rows = _papers(ids)
    return sorted(rows, key=lambda row: (-int(row["year"] or 0), str(row["title"])))


def cited_by(paper_id: int) -> list[dict[str, Any]]:
    """What cites this paper, newest first."""
    engine = connect()
    with Session(engine) as session:
        ids = [
            int(citing)
            for citing in session.exec(
                select(Citation.citing_id).where(Citation.cited_id == paper_id)
            )
            if citing is not None
        ]
    rows = _papers(ids)
    return sorted(rows, key=lambda row: (-int(row["year"] or 0), str(row["title"])))


# --- chunks and FTS ---------------------------------------------------------


def _fts_query(query: str) -> str:
    """Wrap the words in double quotes, so FTS5 syntax characters cannot raise."""
    return " ".join(f'"{word}"' for word in re.findall(r"\w+", query))


def replace_chunks(session: Session, paper_id: int, pages: Sequence[str]) -> int:
    """Replace a paper's text chunks and keep the FTS index in step."""
    connection = session.connection()
    old = [
        int(chunk_id)
        for chunk_id in session.exec(
            select(Chunk.id).where(cast("Any", Chunk.paper_id) == paper_id)
        )
        if chunk_id is not None
    ]
    if old:
        placeholders = ",".join(str(chunk_id) for chunk_id in old)
        connection.exec_driver_sql(f"DELETE FROM chunk_fts WHERE rowid IN ({placeholders})")
        session.exec(delete(Chunk).where(cast("Any", Chunk.paper_id) == paper_id))
        session.flush()
    written = 0
    for page_number, page_text in enumerate(pages):
        page_text = page_text.strip()
        if not page_text:
            continue
        chunk = Chunk(paper_id=paper_id, page=page_number, text=page_text)
        session.add(chunk)
        session.flush()
        connection.exec_driver_sql(
            "INSERT INTO chunk_fts(rowid, text) VALUES (?, ?)", (chunk.id, chunk.text)
        )
        written += 1
    return written


def fts_search(query: str, limit: int) -> list[int]:
    """Chunk ids for a keyword query, best first (bm25 is smaller-is-better)."""
    clauses = _fts_query(query)
    if not clauses:
        return []
    engine = connect()
    with engine.connect() as connection:
        rows = connection.exec_driver_sql(
            "SELECT rowid FROM chunk_fts WHERE chunk_fts MATCH ? ORDER BY bm25(chunk_fts) LIMIT ?",
            (clauses, limit),
        ).fetchall()
    return [int(row[0]) for row in rows]


def chunk_context(chunk_ids: Sequence[int]) -> dict[int, dict[str, Any]]:
    """{chunk_id: {title, page, text, paper_id}} for the chunks a retrieval step picked."""
    if not chunk_ids:
        return {}
    engine = connect()
    with Session(engine) as session:
        statement = (
            select(Chunk)
            .where(cast("Any", Chunk.id).in_(chunk_ids))
            .options(selectinload(cast("Any", Chunk.paper)))
        )
        return {
            int(chunk.id): {
                "paper_id": chunk.paper_id,
                "title": chunk.paper.title if chunk.paper else "",
                "page": chunk.page,
                "text": chunk.text,
            }
            for chunk in session.exec(statement)
            if chunk.id is not None
        }
